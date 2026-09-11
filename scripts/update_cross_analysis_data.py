"""公的統計を都道府県単位で結合し、宿泊DXラボの分析JSONを生成する。"""

from __future__ import annotations

import argparse
import json
import math
import re
from datetime import date
from io import BytesIO
from pathlib import Path
from typing import Any, Iterable

import requests
from openpyxl import load_workbook

try:
    from .update_inbound_data import (
        PREFECTURES,
        SOURCE_PAGE as LODGING_SOURCE_PAGE,
        discover_releases,
        normalize_text,
        serialize_json,
        write_if_changed,
    )
except ImportError:  # `python scripts/update_cross_analysis_data.py` での実行用
    from update_inbound_data import (
        PREFECTURES,
        SOURCE_PAGE as LODGING_SOURCE_PAGE,
        discover_releases,
        normalize_text,
        serialize_json,
        write_if_changed,
    )


POPULATION_CURRENT_URL = "https://www.stat.go.jp/data/jinsui/2024np/zuhyou/05k2024-3.xlsx"
POPULATION_PREVIOUS_URL = "https://www.stat.go.jp/data/jinsui/2019np/zuhyou/05k01-3.xlsx"
POPULATION_SOURCE_URL = "https://www.stat.go.jp/data/jinsui/2024np/index.htm"
ECONOMIC_CENSUS_URL = (
    "https://www.e-stat.go.jp/stat-search/file-download?"
    "fileKind=0&statInfId=000040389322"
)
ECONOMIC_CENSUS_SOURCE_URL = "https://www.e-stat.go.jp/dbview?sid=0004040078"
ECONOMIC_CENSUS_STATS_DATA_ID = "0004040078"
ECONOMIC_CENSUS_STAT_INF_ID = "000040389322"
USER_AGENT = "ShukuhakuDXLab-CrossAnalysisUpdater/1.0 (+https://lab.ugatta-llc.com/)"

PREFECTURE_SLUGS = [
    "hokkaido", "aomori", "iwate", "miyagi", "akita", "yamagata", "fukushima",
    "ibaraki", "tochigi", "gunma", "saitama", "chiba", "tokyo", "kanagawa",
    "niigata", "toyama", "ishikawa", "fukui", "yamanashi", "nagano", "gifu",
    "shizuoka", "aichi", "mie", "shiga", "kyoto", "osaka", "hyogo", "nara",
    "wakayama", "tottori", "shimane", "okayama", "hiroshima", "yamaguchi",
    "tokushima", "kagawa", "ehime", "kochi", "fukuoka", "saga", "nagasaki",
    "kumamoto", "oita", "miyazaki", "kagoshima", "okinawa",
]
PREFECTURE_BY_CODE = {
    f"{index:02d}": {"name": name, "slug": slug}
    for index, (name, slug) in enumerate(zip(PREFECTURES, PREFECTURE_SLUGS), 1)
}


def fetch_bytes(session: requests.Session, url: str) -> bytes:
    response = session.get(url, timeout=120)
    response.raise_for_status()
    if len(response.content) < 10_000:
        raise RuntimeError(f"統計ファイルが小さすぎます: {url}")
    return response.content


def find_sheet(workbook: Any, title_prefix: str, required: Iterable[str]) -> Any:
    candidates = []
    for sheet in workbook.worksheets:
        first = normalize_text(sheet.cell(1, 1).value)
        if first.startswith(title_prefix) and all(term in first for term in required):
            candidates.append(sheet)
    if len(candidates) != 1:
        raise RuntimeError(f"統計表を一意に特定できません: {title_prefix} / {[s.title for s in candidates]}")
    return candidates[0]


def prefecture_code_from_label(value: Any) -> str | None:
    match = re.match(r"^\s*(\d{2})", normalize_text(value))
    return match.group(1) if match and match.group(1) in PREFECTURE_BY_CODE else None


def find_header_column(sheet: Any, needle: str, *, exclude: str | None = None) -> int:
    candidates: list[int] = []
    for row in range(3, min(sheet.max_row, 12) + 1):
        for column in range(2, sheet.max_column + 1):
            text = normalize_text(sheet.cell(row, column).value).replace(" ", "")
            compact_needle = needle.replace(" ", "")
            compact_exclude = exclude.replace(" ", "") if exclude else None
            if compact_needle in text and (compact_exclude is None or compact_exclude not in text):
                candidates.append(column)
    if not candidates:
        raise RuntimeError(f"統計表の列見出しを特定できません: {needle}")
    return min(candidates)


def number_or_none(value: Any) -> float | None:
    if value is None or normalize_text(value) in {"", "-", "―", "—", "...", "X"}:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)) and math.isfinite(float(value)):
        return float(value)
    text = normalize_text(value).replace(",", "")
    return float(text) if re.fullmatch(r"-?\d+(?:\.\d+)?", text) else None


def extract_prefecture_rows(sheet: Any, value_columns: tuple[int, ...]) -> tuple[list[float | None], dict[str, list[float | None]]]:
    national: list[float | None] | None = None
    prefectures: dict[str, list[float | None]] = {}
    for row in sheet.iter_rows(min_row=1, values_only=True):
        label = row[0] if row else None
        code = prefecture_code_from_label(label)
        values = [number_or_none(row[index - 1] if len(row) >= index else None) for index in value_columns]
        if code:
            prefectures[code] = values
        elif national is None and re.search(r"(?:令和\d+年|20\d{2}年).*?\d{1,2}月", normalize_text(label)):
            national = values
    if national is None or set(prefectures) != set(PREFECTURE_BY_CODE):
        missing = sorted(set(PREFECTURE_BY_CODE) - set(prefectures))
        raise RuntimeError(f"宿泊統計の全国値または都道府県が不足しています: {missing}")
    return national, prefectures


def load_lodging_month(session: requests.Session, release: Any) -> dict[str, Any]:
    workbook = load_workbook(BytesIO(fetch_bytes(session, release.url)), read_only=True, data_only=True)
    try:
        guest_sheet = find_sheet(workbook, "第2表", ("延べ宿泊者数",))
        rooms_sheet = find_sheet(workbook, "第7表", ("利用客室数",))
        occupancy_sheet = find_sheet(workbook, "第8表", ("客室稼働率",))
        total_guest_column = find_header_column(guest_sheet, "延べ宿泊者数", exclude="外国人")
        foreign_guest_column = find_header_column(guest_sheet, "外国人延べ宿泊者数")
        used_rooms_column = find_header_column(rooms_sheet, "利用客室数")
        occupancy_column = find_header_column(occupancy_sheet, "客室稼働率")
        national_guest, prefecture_guest = extract_prefecture_rows(guest_sheet, (total_guest_column, foreign_guest_column))
        national_rooms, prefecture_rooms = extract_prefecture_rows(rooms_sheet, (used_rooms_column,))
        national_occupancy, prefecture_occupancy = extract_prefecture_rows(occupancy_sheet, (occupancy_column,))
    finally:
        workbook.close()

    def combine(guest: list[float | None], rooms: list[float | None], occupancy: list[float | None]) -> dict[str, Any]:
        return {
            "total_guest_nights": int(guest[0]) if guest[0] is not None else None,
            "foreign_guest_nights": int(guest[1]) if guest[1] is not None else None,
            "used_rooms": int(rooms[0]) if rooms[0] is not None else None,
            "occupancy_rate": occupancy[0],
        }

    return {
        "period": f"{release.year:04d}-{release.month:02d}",
        "release_type": "第2次速報",
        "source_url": release.url,
        "national": combine(national_guest, national_rooms, national_occupancy),
        "prefectures": {
            code: combine(prefecture_guest[code], prefecture_rooms[code], prefecture_occupancy[code])
            for code in PREFECTURE_BY_CODE
        },
    }


def annualize_lodging(records: list[dict[str, Any]]) -> dict[str, Any]:
    total_guest_nights = sum(item["total_guest_nights"] for item in records if item["total_guest_nights"] is not None)
    foreign_guest_nights = sum(item["foreign_guest_nights"] for item in records if item["foreign_guest_nights"] is not None)
    if sum(item["total_guest_nights"] is not None for item in records) != len(records):
        total_guest_nights = None
    if sum(item["foreign_guest_nights"] is not None for item in records) != len(records):
        foreign_guest_nights = None
    used_rooms = 0.0
    available_rooms = 0.0
    for item in records:
        used = item["used_rooms"]
        occupancy = item["occupancy_rate"]
        if used is None or occupancy is None or occupancy <= 0:
            continue
        used_rooms += used
        available_rooms += used / (occupancy / 100)
    occupancy_rate = round(used_rooms / available_rooms * 100, 1) if available_rooms > 0 else None
    return {
        "total_guest_nights": total_guest_nights,
        "foreign_guest_nights": foreign_guest_nights,
        "occupancy_rate": occupancy_rate,
    }


def load_lodging_stats(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    releases, source_updated = discover_releases(session)
    if len(releases) < 24:
        raise RuntimeError("同じ第2次速報で24か月分を取得できません。")
    selected = sorted(releases[:24], key=lambda item: (item.year, item.month))
    month_numbers = [release.year * 12 + release.month for release in selected]
    if any(current - previous != 1 for previous, current in zip(month_numbers, month_numbers[1:])):
        raise RuntimeError("第2次速報の24か月が連続していません。")
    months = [load_lodging_month(session, release) for release in selected]
    current_period = months[-12:]
    previous_period = months[:12]

    def aggregate(key: str) -> dict[str, Any]:
        current = annualize_lodging([item[key] for item in current_period])
        previous = annualize_lodging([item[key] for item in previous_period])
        growth = None
        if current["total_guest_nights"] is not None and previous["total_guest_nights"]:
            growth = current["total_guest_nights"] / previous["total_guest_nights"] - 1
        return {
            "total_guest_nights": current["total_guest_nights"],
            "total_guest_nights_previous": previous["total_guest_nights"],
            "foreign_guest_nights": current["foreign_guest_nights"],
            "occupancy_rate": current["occupancy_rate"],
            "demand_growth_rate": growth,
        }

    data = {
        "national": aggregate("national"),
        "prefectures": {
            code: aggregate_prefecture(code, current_period, previous_period)
            for code in PREFECTURE_BY_CODE
        },
    }
    metadata = {
        "source_name": "観光庁「宿泊旅行統計調査」",
        "table_name": "第2表（延べ宿泊者数）・第7表（利用客室数）・第8表（客室稼働率）",
        "statistics_code": "00601020",
        "reference_period": f"{current_period[0]['period']}〜{current_period[-1]['period']}",
        "comparison_period": f"{previous_period[0]['period']}〜{previous_period[-1]['period']}",
        "release_type": "第2次速報",
        "published_date": source_updated,
        "source_url": LODGING_SOURCE_PAGE,
        "latest_excel_url": months[-1]["source_url"],
    }
    return data, metadata


def aggregate_prefecture(code: str, current_period: list[dict[str, Any]], previous_period: list[dict[str, Any]]) -> dict[str, Any]:
    current = annualize_lodging([item["prefectures"][code] for item in current_period])
    previous = annualize_lodging([item["prefectures"][code] for item in previous_period])
    growth = None
    if current["total_guest_nights"] is not None and previous["total_guest_nights"]:
        growth = current["total_guest_nights"] / previous["total_guest_nights"] - 1
    return {
        "total_guest_nights": current["total_guest_nights"],
        "total_guest_nights_previous": previous["total_guest_nights"],
        "foreign_guest_nights": current["foreign_guest_nights"],
        "occupancy_rate": current["occupancy_rate"],
        "demand_growth_rate": growth,
    }


def load_population_table(content: bytes, expected_year: int) -> dict[str, dict[str, int]]:
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        title = normalize_text(sheet.cell(1, 1).value)
        headers = " ".join(normalize_text(sheet.cell(row, col).value) for row in range(1, 13) for col in range(1, 17))
        if str(expected_year) not in title or "15~64歳" not in headers.replace("～", "~"):
            raise RuntimeError(f"人口推計の表構造を確認できません: {expected_year}")
        records: dict[str, dict[str, int]] = {}
        for row in sheet.iter_rows(min_row=13, values_only=True):
            code = normalize_text(row[1] if len(row) > 1 else None)
            if code not in PREFECTURE_BY_CODE:
                continue
            name = normalize_text(row[2] if len(row) > 2 else None).replace(" ", "")
            if name != PREFECTURE_BY_CODE[code]["name"]:
                raise RuntimeError(f"人口推計の都道府県コードと名称が一致しません: {code} {name}")
            age_groups = [number_or_none(row[index]) for index in (4, 5, 6)]
            if any(value is None for value in age_groups):
                raise RuntimeError(f"人口推計の年齢3区分が不足しています: {name}")
            records[code] = {
                "population": int(sum(age_groups) * 1000),
                "working_age_population": int(age_groups[1] * 1000),
            }
        if set(records) != set(PREFECTURE_BY_CODE):
            raise RuntimeError("人口推計に47都道府県が揃っていません。")
        return records
    finally:
        workbook.close()


def load_population_stats(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    current = load_population_table(fetch_bytes(session, POPULATION_CURRENT_URL), 2024)
    previous = load_population_table(fetch_bytes(session, POPULATION_PREVIOUS_URL), 2019)
    result = {}
    for code in PREFECTURE_BY_CODE:
        current_working = current[code]["working_age_population"]
        previous_working = previous[code]["working_age_population"]
        result[code] = {
            "population": current[code]["population"],
            "working_age_population": current_working,
            "working_age_population_previous": previous_working,
            "working_age_population_change_rate": current_working / previous_working - 1,
        }
    metadata = {
        "source_name": "総務省統計局「人口推計」",
        "table_name": "第3表 都道府県、年齢（3区分）、男女別人口―総人口",
        "statistics_code": "00200524",
        "reference_period": "2024-10-01",
        "comparison_period": "2019-10-01",
        "published_date": "2025-04-14",
        "source_url": POPULATION_SOURCE_URL,
        "current_excel_url": POPULATION_CURRENT_URL,
        "previous_excel_url": POPULATION_PREVIOUS_URL,
        "unit_note": "千人単位の公表値を人へ換算。年齢3区分の合計を総人口として使用。",
    }
    return result, metadata


def load_economic_census(session: requests.Session) -> tuple[dict[str, Any], dict[str, Any]]:
    workbook = load_workbook(BytesIO(fetch_bytes(session, ECONOMIC_CENSUS_URL)), read_only=True, data_only=True)
    try:
        sheet = workbook.active
        header = [normalize_text(sheet.cell(7, col).value) for col in range(1, 10)]
        if not {"地域区分", "産業中分類", "経営組織"}.issubset(set(header)):
            raise RuntimeError("経済センサス表2の見出しを確認できません。")
        result: dict[str, dict[str, int]] = {}
        for row in sheet.iter_rows(min_row=8, values_only=True):
            region = normalize_text(row[1] if len(row) > 1 else None)
            match = re.match(r"(\d{2})000_(.+)", region)
            if not match or match.group(1) not in PREFECTURE_BY_CODE:
                continue
            code = match.group(1)
            name = match.group(2)
            if name != PREFECTURE_BY_CODE[code]["name"]:
                continue
            industry = normalize_text(row[4] if len(row) > 4 else None)
            organization = normalize_text(row[5] if len(row) > 5 else None)
            if organization != "0_総数":
                continue
            establishments = number_or_none(row[6] if len(row) > 6 else None)
            employees = number_or_none(row[7] if len(row) > 7 else None)
            if establishments is None or employees is None:
                continue
            item = result.setdefault(code, {})
            if industry.startswith("75_宿泊業"):
                item["lodging_establishments"] = int(establishments)
                item["lodging_employees"] = int(employees)
            elif industry.startswith("AR_全産業"):
                item["all_establishments"] = int(establishments)
                item["all_employees"] = int(employees)
        required = {"lodging_establishments", "lodging_employees", "all_establishments", "all_employees"}
        if set(result) != set(PREFECTURE_BY_CODE) or any(not required <= set(item) for item in result.values()):
            raise RuntimeError("経済センサスに必要な47都道府県データが揃っていません。")
    finally:
        workbook.close()
    metadata = {
        "source_name": "総務省・経済産業省「令和6年経済センサス―基礎調査」（e-Stat）",
        "table_name": "甲調査（民営事業所）事業所に関する集計 表2 産業（中分類）×経営組織別",
        "reference_period": "2024-06",
        "published_date": "2025-12-24",
        "source_url": ECONOMIC_CENSUS_SOURCE_URL,
        "excel_url": ECONOMIC_CENSUS_URL,
        "stats_data_id": ECONOMIC_CENSUS_STATS_DATA_ID,
        "stat_inf_id": ECONOMIC_CENSUS_STAT_INF_ID,
        "scope_note": "民営事業所（雇用者のいない個人経営の事業所を除く）。全産業は公務を除く。",
    }
    return result, metadata


def safe_ratio(numerator: float | int | None, denominator: float | int | None) -> float | None:
    if numerator is None or denominator is None or denominator == 0:
        return None
    return numerator / denominator


def percentile_rank(values: dict[str, float | None]) -> dict[str, float | None]:
    valid = sorted((value, code) for code, value in values.items() if value is not None and math.isfinite(value))
    result: dict[str, float | None] = {code: None for code in values}
    count = len(valid)
    if count == 1:
        result[valid[0][1]] = 50.0
        return result
    index = 0
    while index < count:
        end = index
        while end + 1 < count and valid[end + 1][0] == valid[index][0]:
            end += 1
        score = ((index + end) / 2) / (count - 1) * 100
        for position in range(index, end + 1):
            result[valid[position][1]] = round(score, 1)
        index = end + 1
    return result


def calculate_ranking(values: dict[str, float | None]) -> dict[str, int | None]:
    valid = [value for value in values.values() if value is not None]
    return {
        code: 1 + sum(other > value for other in valid) if value is not None else None
        for code, value in values.items()
    }


def weighted_score(components: list[tuple[float | None, float]]) -> int | None:
    if any(value is None for value, _ in components):
        return None
    return round(sum(float(value) * weight for value, weight in components))


def score_level(score: int | None) -> str | None:
    if score is None:
        return None
    if score >= 80:
        return "非常に高い"
    if score >= 60:
        return "高い"
    if score >= 40:
        return "中程度"
    if score >= 20:
        return "低い"
    return "比較的低い"


def classify_supply_market(growth: float, occupancy: float, national_growth: float, national_occupancy: float) -> str:
    return {
        (True, True): "需給ひっ迫候補",
        (True, False): "需要成長・供給余力型",
        (False, True): "高稼働・成熟型",
        (False, False): "需給軟調型",
    }[(growth > national_growth, occupancy > national_occupancy)]


def classify_tourism_region(pressure: int, dependency: int) -> str:
    return {
        (True, True): "観光重要度・負荷ともに高い地域",
        (False, True): "観光産業比重型",
        (True, False): "来訪集中型",
        (False, False): "分散型",
    }[(pressure >= 50, dependency >= 50)]


def build_datasets(session: requests.Session, previous_common: dict[str, Any] | None = None) -> dict[str, dict[str, Any]]:
    lodging, lodging_source = load_lodging_stats(session)
    population, population_source = load_population_stats(session)
    census, census_source = load_economic_census(session)
    retrieved_at = date.today().isoformat()
    previous_sources = (previous_common or {}).get("metadata", {}).get("sources", {})
    current_signature = (
        lodging_source["reference_period"], population_source["reference_period"], census_source["reference_period"]
    )
    previous_signature = tuple(
        previous_sources.get(key, {}).get("reference_period")
        for key in ("lodging", "population", "economic_census")
    )
    if current_signature == previous_signature:
        retrieved_at = (previous_common or {}).get("metadata", {}).get("retrieved_at", retrieved_at)
    for source in (lodging_source, population_source, census_source):
        source["retrieved_at"] = retrieved_at

    prefectures: dict[str, Any] = {}
    for code, identity in PREFECTURE_BY_CODE.items():
        item = {"prefecture_code": code, "prefecture_slug": identity["slug"], "prefecture_name": identity["name"]}
        item.update(population[code])
        item.update(lodging["prefectures"][code])
        item.update(census[code])
        item.update({
            "data_period": lodging_source["reference_period"],
            "source_dates": {
                "lodging": lodging_source["reference_period"],
                "population": population_source["reference_period"],
                "economic_census": census_source["reference_period"],
            },
            "lodging_demand_per_employee": safe_ratio(item["total_guest_nights"], item["lodging_employees"]),
            "lodging_demand_per_establishment": safe_ratio(item["total_guest_nights"], item["lodging_establishments"]),
            "lodging_density": safe_ratio(item["total_guest_nights"], item["population"]),
            "foreign_lodging_density": safe_ratio(item["foreign_guest_nights"], item["population"]),
            "lodging_establishments_per_10000": safe_ratio(item["lodging_establishments"] * 10_000, item["population"]),
            "lodging_employee_share": safe_ratio(item["lodging_employees"], item["all_employees"]),
            "lodging_establishment_share": safe_ratio(item["lodging_establishments"], item["all_establishments"]),
        })
        prefectures[code] = item

    national = dict(lodging["national"])
    national_population = sum(item["population"] for item in population.values())
    national.update({
        "population": national_population,
        "working_age_population": sum(item["working_age_population"] for item in population.values()),
        "working_age_population_previous": sum(item["working_age_population_previous"] for item in population.values()),
        "lodging_establishments": sum(item["lodging_establishments"] for item in census.values()),
        "lodging_employees": sum(item["lodging_employees"] for item in census.values()),
        "all_establishments": sum(item["all_establishments"] for item in census.values()),
        "all_employees": sum(item["all_employees"] for item in census.values()),
    })
    national["working_age_population_change_rate"] = safe_ratio(national["working_age_population"], national["working_age_population_previous"])
    if national["working_age_population_change_rate"] is not None:
        national["working_age_population_change_rate"] -= 1
    national.update({
        "lodging_demand_per_employee": safe_ratio(national["total_guest_nights"], national["lodging_employees"]),
        "lodging_demand_per_establishment": safe_ratio(national["total_guest_nights"], national["lodging_establishments"]),
        "lodging_density": safe_ratio(national["total_guest_nights"], national["population"]),
        "foreign_lodging_density": safe_ratio(national["foreign_guest_nights"], national["population"]),
        "lodging_establishments_per_10000": safe_ratio(national["lodging_establishments"] * 10_000, national["population"]),
        "lodging_employee_share": safe_ratio(national["lodging_employees"], national["all_employees"]),
        "lodging_establishment_share": safe_ratio(national["lodging_establishments"], national["all_establishments"]),
    })

    common_metadata = {
        "generated_by": "宿泊DXラボ",
        "retrieved_at": retrieved_at,
        "prefecture_count": 47,
        "sources": {"lodging": lodging_source, "population": population_source, "economic_census": census_source},
        "methodology_note": "宿泊旅行統計調査は2026年1月から層化基準が変更されており、2025年以前との比較には変更の影響が含まれる可能性があります。",
    }
    common = {"metadata": common_metadata, "national": national, "prefectures": prefectures}

    growth_p = percentile_rank({code: item["demand_growth_rate"] for code, item in prefectures.items()})
    occupancy_p = percentile_rank({code: item["occupancy_rate"] for code, item in prefectures.items()})
    employee_load_p = percentile_rank({code: item["lodging_demand_per_employee"] for code, item in prefectures.items()})
    workforce_decline_p = percentile_rank({code: -item["working_age_population_change_rate"] for code, item in prefectures.items()})
    dx_records = {}
    for code, item in prefectures.items():
        score = weighted_score([(growth_p[code], .20), (occupancy_p[code], .25), (employee_load_p[code], .30), (workforce_decline_p[code], .25)])
        dx_records[code] = {**item, "percentiles": {"demand_growth": growth_p[code], "occupancy": occupancy_p[code], "lodging_operation_load": employee_load_p[code], "workforce_decline": workforce_decline_p[code]}, "score": score, "score_level": score_level(score)}
    dx_ranks = calculate_ranking({code: item["score"] for code, item in dx_records.items()})
    for code in dx_records:
        dx_records[code]["rank"] = dx_ranks[code]

    establishment_load_p = percentile_rank({code: item["lodging_demand_per_establishment"] for code, item in prefectures.items()})
    supply_records = {}
    for code, item in prefectures.items():
        score = weighted_score([(occupancy_p[code], .40), (growth_p[code], .35), (establishment_load_p[code], .25)])
        market_type = classify_supply_market(item["demand_growth_rate"], item["occupancy_rate"], national["demand_growth_rate"], national["occupancy_rate"])
        supply_records[code] = {**item, "percentiles": {"demand_growth": growth_p[code], "occupancy": occupancy_p[code], "demand_per_establishment": establishment_load_p[code]}, "score": score, "market_type": market_type}
    supply_ranks = calculate_ranking({code: item["score"] for code, item in supply_records.items()})
    for code in supply_records:
        supply_records[code]["rank"] = supply_ranks[code]

    lodging_density_p = percentile_rank({code: item["lodging_density"] for code, item in prefectures.items()})
    establishment_density_p = percentile_rank({code: item["lodging_establishments_per_10000"] for code, item in prefectures.items()})
    employee_share_p = percentile_rank({code: item["lodging_employee_share"] for code, item in prefectures.items()})
    establishment_share_p = percentile_rank({code: item["lodging_establishment_share"] for code, item in prefectures.items()})
    tourism_records = {}
    for code, item in prefectures.items():
        pressure = weighted_score([(lodging_density_p[code], .70), (establishment_density_p[code], .30)])
        dependency = weighted_score([(employee_share_p[code], .50), (establishment_share_p[code], .25), (lodging_density_p[code], .25)])
        region_type = classify_tourism_region(pressure, dependency)
        tourism_records[code] = {**item, "percentiles": {"lodging_density": lodging_density_p[code], "establishment_density": establishment_density_p[code], "employee_share": employee_share_p[code], "establishment_share": establishment_share_p[code]}, "pressure_score": pressure, "dependency_score": dependency, "region_type": region_type}
    pressure_ranks = calculate_ranking({code: item["pressure_score"] for code, item in tourism_records.items()})
    dependency_ranks = calculate_ranking({code: item["dependency_score"] for code, item in tourism_records.items()})
    for code in tourism_records:
        tourism_records[code]["pressure_rank"] = pressure_ranks[code]
        tourism_records[code]["dependency_rank"] = dependency_ranks[code]

    metadata = {
        "common_data": "prefecture-statistics.json",
        "retrieved_at": retrieved_at,
        "sources": common_metadata["sources"],
        "methodology_note": common_metadata["methodology_note"],
        "indicator_note": "公的統計をもとに宿泊DXラボが独自に算出した参考指標です。",
    }
    return {
        "prefecture-statistics.json": common,
        "dx-necessity.json": {"metadata": {**metadata, "formula": "需要伸び率20%＋客室稼働率25%＋宿泊運営負荷30%＋生産年齢人口減少25%（各percentile rank）"}, "national": national, "prefectures": dx_records},
        "supply-demand-gap.json": {"metadata": {**metadata, "formula": "客室稼働率40%＋需要伸び率35%＋事業所あたり宿泊需要25%（各percentile rank）", "quadrant_basis": "全国の需要伸び率と客室稼働率"}, "national": national, "prefectures": supply_records},
        "tourism-pressure.json": {"metadata": {**metadata, "pressure_formula": "宿泊密度70%＋宿泊施設密度30%（各percentile rank）", "dependency_formula": "宿泊業従業者比率50%＋宿泊業事業所比率25%＋宿泊密度25%（各percentile rank）", "type_basis": "各参考スコア50点を相対比較の境界として分類"}, "national": national, "prefectures": tourism_records},
    }


def ensure_finite(value: Any, path: str = "root") -> None:
    if isinstance(value, dict):
        for key, child in value.items():
            ensure_finite(child, f"{path}.{key}")
    elif isinstance(value, list):
        for index, child in enumerate(value):
            ensure_finite(child, f"{path}[{index}]")
    elif isinstance(value, float) and not math.isfinite(value):
        raise ValueError(f"NaNまたはInfinityを検出しました: {path}")


def validate_datasets(datasets: dict[str, dict[str, Any]]) -> None:
    expected = set(PREFECTURE_BY_CODE)
    for filename, dataset in datasets.items():
        records = dataset.get("prefectures", {})
        if set(records) != expected or len(records) != 47:
            raise ValueError(f"{filename}: 47都道府県が揃っていません。")
        if len({item["prefecture_code"] for item in records.values()}) != 47:
            raise ValueError(f"{filename}: 都道府県コードが重複しています。")
        for code, item in records.items():
            if item["population"] <= 0:
                raise ValueError(f"{filename}: {code}の人口が不正です。")
            for key in ("total_guest_nights", "foreign_guest_nights", "lodging_establishments", "lodging_employees", "all_establishments", "all_employees"):
                if item[key] is not None and item[key] < 0:
                    raise ValueError(f"{filename}: {code}.{key}が不正です。")
            if item["occupancy_rate"] is not None and not 0 <= item["occupancy_rate"] <= 100:
                raise ValueError(f"{filename}: {code}の稼働率が不正です。")
            for key in ("score", "pressure_score", "dependency_score"):
                if key in item and item[key] is not None and not 0 <= item[key] <= 100:
                    raise ValueError(f"{filename}: {code}.{key}が範囲外です。")
        ensure_finite(dataset)


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--output-dir", type=Path, default=Path("data/analysis"))
    args = parser.parse_args()
    previous_common = None
    common_path = args.output_dir / "prefecture-statistics.json"
    if common_path.exists():
        previous_common = json.loads(common_path.read_text(encoding="utf-8"))
    session = requests.Session()
    session.headers.update({"User-Agent": USER_AGENT})
    datasets = build_datasets(session, previous_common)
    validate_datasets(datasets)
    changed = False
    for filename, dataset in datasets.items():
        changed = write_if_changed(args.output_dir / filename, serialize_json(dataset)) or changed
    print("分析JSONを更新しました。" if changed else "新しい変更はありません。")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
